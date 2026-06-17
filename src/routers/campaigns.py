import json
import csv
import io
from datetime import datetime, timedelta

from fastapi import APIRouter, File, HTTPException, Query, UploadFile

from src.api_helpers import *
from src.email_verify import verify_email
from src.models import EnrichmentMode, Lead, LeadStatus, PipelineRun, RunStatus, Segment


router = APIRouter()


def _pct(numerator: int, denominator: int) -> float:
    if not denominator:
        return 0.0
    return round((float(numerator) / float(denominator)) * 100, 2)


def _report_cutoff(days: int) -> tuple[str, list[str]]:
    normalized_days = max(1, min(int(days or 30), 365))
    today = datetime.utcnow().date()
    start_date = today - timedelta(days=normalized_days - 1)
    cutoff = datetime.combine(start_date, datetime.min.time()).strftime("%Y-%m-%d %H:%M:%S")

    dates = [
        (start_date + timedelta(days=index)).isoformat()
        for index in range(normalized_days)
    ]

    return cutoff, dates


def _report_touch_name_map(campaign_filename: str) -> dict[int, str]:
    names: dict[int, str] = {}

    for step in _active_steps(campaign_filename):
        names[int(step.touch_number or 0)] = (
            step.touch_name or f"Touch {step.touch_number}"
        )

    return names


def _campaign_report_data(campaign_filename: str, days: int = 30) -> dict:
    cutoff, date_keys = _report_cutoff(days)
    conn = lead_repo.db.conn()

    touch_names = _report_touch_name_map(campaign_filename)

    sent_rows = conn.execute(
        """
        SELECT COALESCE(touch_number, 0) AS touch_number,
               COUNT(*) AS sent
        FROM send_log
        WHERE campaign_filename = ?
          AND sent_at >= ?
        GROUP BY COALESCE(touch_number, 0)
        """,
        (campaign_filename, cutoff),
    ).fetchall()

    sent_by_touch = {
        int(row["touch_number"] or 0): int(row["sent"] or 0)
        for row in sent_rows
    }

    reply_rows = conn.execute(
        """
        SELECT touch_number, COUNT(*) AS replies
        FROM (
            SELECT a.lead_id,
                   COALESCE((
                       SELECT MAX(s.touch_number)
                       FROM send_log s
                       WHERE s.lead_id = a.lead_id
                         AND s.campaign_filename = a.campaign_filename
                         AND s.sent_at <= a.created_at
                   ), 0) AS touch_number
            FROM lead_activities a
            WHERE a.campaign_filename = ?
              AND a.activity_type IN ('reply_detected', 'replied')
              AND a.created_at >= ?
            GROUP BY a.lead_id
        )
        GROUP BY touch_number
        """,
        (campaign_filename, cutoff),
    ).fetchall()

    replies_by_touch = {
        int(row["touch_number"] or 0): int(row["replies"] or 0)
        for row in reply_rows
    }

    bounce_rows = conn.execute(
        """
        SELECT touch_number, COUNT(*) AS bounces
        FROM (
            SELECT a.lead_id,
                   COALESCE((
                       SELECT MAX(s.touch_number)
                       FROM send_log s
                       WHERE s.lead_id = a.lead_id
                         AND s.campaign_filename = a.campaign_filename
                         AND s.sent_at <= a.created_at
                   ), 0) AS touch_number
            FROM lead_activities a
            WHERE a.campaign_filename = ?
              AND a.activity_type IN ('bounce_detected', 'bounced')
              AND a.created_at >= ?
            GROUP BY a.lead_id
        )
        GROUP BY touch_number
        """,
        (campaign_filename, cutoff),
    ).fetchall()

    bounces_by_touch = {
        int(row["touch_number"] or 0): int(row["bounces"] or 0)
        for row in bounce_rows
    }

    all_touch_numbers = sorted(
        set(touch_names.keys())
        | set(sent_by_touch.keys())
        | set(replies_by_touch.keys())
        | set(bounces_by_touch.keys())
    )

    per_touch = []
    for touch_number in all_touch_numbers:
        if touch_number <= 0:
            continue

        sent = sent_by_touch.get(touch_number, 0)
        replies = replies_by_touch.get(touch_number, 0)
        bounces = bounces_by_touch.get(touch_number, 0)

        per_touch.append(
            {
                "touch_number": touch_number,
                "name": touch_names.get(touch_number) or f"Touch {touch_number}",
                "sent": sent,
                "replies_attributed": replies,
                "bounces": bounces,
                "reply_rate": _pct(replies, sent),
            }
        )

    daily_sent_rows = conn.execute(
        """
        SELECT substr(sent_at, 1, 10) AS date,
               COUNT(*) AS sent
        FROM send_log
        WHERE campaign_filename = ?
          AND sent_at >= ?
        GROUP BY substr(sent_at, 1, 10)
        ORDER BY date
        """,
        (campaign_filename, cutoff),
    ).fetchall()

    daily_replies_rows = conn.execute(
        """
        SELECT substr(created_at, 1, 10) AS date,
               COUNT(DISTINCT lead_id) AS replies
        FROM lead_activities
        WHERE campaign_filename = ?
          AND activity_type IN ('reply_detected', 'replied')
          AND created_at >= ?
        GROUP BY substr(created_at, 1, 10)
        ORDER BY date
        """,
        (campaign_filename, cutoff),
    ).fetchall()

    sent_by_day = {
        row["date"]: int(row["sent"] or 0)
        for row in daily_sent_rows
    }
    replies_by_day = {
        row["date"]: int(row["replies"] or 0)
        for row in daily_replies_rows
    }

    daily = [
        {
            "date": date_key,
            "sent": sent_by_day.get(date_key, 0),
            "replies": replies_by_day.get(date_key, 0),
        }
        for date_key in date_keys
    ]

    status_rows = conn.execute(
        """
        SELECT status, COUNT(*) AS total
        FROM lead_sequence_state
        WHERE campaign_filename = ?
        GROUP BY status
        """,
        (campaign_filename,),
    ).fetchall()

    status_breakdown = {
        "active": 0,
        "not_started": 0,
        "draft_generated": 0,
        "enriched": 0,
        "waiting_followup": 0,
        "completed": 0,
        "replied": 0,
        "bounced": 0,
        "unsubscribed": 0,
    }

    for row in status_rows:
        status = row["status"] or "unknown"
        total = int(row["total"] or 0)
        status_breakdown[status] = total

        if status in {
            "not_started",
            "enriched",
            "draft_generated",
            "approved",
            "scheduled",
            "waiting_followup",
        }:
            status_breakdown["active"] += total

    draft_status_rows = conn.execute(
        """
        SELECT status, COUNT(*) AS total
        FROM outreach_drafts
        WHERE campaign_filename = ?
        GROUP BY status
        """,
        (campaign_filename,),
    ).fetchall()

    draft_status_breakdown = {
        "draft": 0,
        "approved": 0,
        "scheduled": 0,
        "sending": 0,
        "sent": 0,
        "failed": 0,
        "skipped": 0,
    }

    for row in draft_status_rows:
        status = row["status"] or "unknown"
        draft_status_breakdown[status] = int(row["total"] or 0)

    total_sent = sum(item["sent"] for item in per_touch)
    total_replies = sum(item["replies_attributed"] for item in per_touch)
    activity_bounces = sum(item["bounces"] for item in per_touch)
    state_bounces = int(status_breakdown.get("bounced", 0) or 0)
    total_bounces = max(activity_bounces, state_bounces)
    total_unsubscribes = int(status_breakdown.get("unsubscribed", 0) or 0)

    return {
        "campaign_filename": campaign_filename,
        "days": max(1, min(int(days or 30), 365)),
        "per_touch": per_touch,
        "totals": {
            "sent": total_sent,
            "drafts": draft_status_breakdown.get("draft", 0)
            + draft_status_breakdown.get("approved", 0),
            "scheduled": draft_status_breakdown.get("scheduled", 0),
            "sending": draft_status_breakdown.get("sending", 0),
            "failed": draft_status_breakdown.get("failed", 0),
            "skipped": draft_status_breakdown.get("skipped", 0),
            "replies": total_replies,
            "bounces": total_bounces,
            "unsubscribes": total_unsubscribes,
            "reply_rate": _pct(total_replies, total_sent),
            "bounce_rate": _pct(total_bounces, total_sent),
        },
        "daily": daily,
        "status_breakdown": status_breakdown,
        "draft_status_breakdown": draft_status_breakdown,
    }


@router.get("/api/dashboard/summary")
def get_dashboard_summary() -> dict:
    policy = SendPolicy().status()
    campaigns = campaign_repo.list_all()
    return {
        "totals": {
            "leads": lead_repo.count_all(),
            "sent_today": send_log_repo.count_today(),
            "todays_cap": policy.get("todays_cap", 0),
            "replies_total": outreach_repo.count_states_by_status("replied"),
            "active_campaigns": len(campaigns),
        },
        "recent_activities": outreach_repo.recent_activities(15),
        "due_today_total": sum(
            len(outreach_repo.due_items(c["filename"], None, None))
            for c in campaigns
        ),
    }




@router.get("/api/campaigns/summary")
def get_campaigns_summary() -> dict:
    campaigns = campaign_repo.list_all()
    filenames = [campaign["filename"] for campaign in campaigns]

    if not filenames:
        return {}

    placeholders = ",".join("?" for _ in filenames)
    conn = lead_repo.db.conn()

    sent_rows = conn.execute(
        f"""
        SELECT campaign_filename, COUNT(*) AS sent
        FROM send_log
        WHERE campaign_filename IN ({placeholders})
        GROUP BY campaign_filename
        """,
        filenames,
    ).fetchall()

    state_rows = conn.execute(
        f"""
        SELECT campaign_filename, status, COUNT(*) AS total
        FROM lead_sequence_state
        WHERE campaign_filename IN ({placeholders})
        GROUP BY campaign_filename, status
        """,
        filenames,
    ).fetchall()

    result = {
        filename: {
            "sent": 0,
            "replies": 0,
            "bounces": 0,
            "unsubscribes": 0,
            "reply_rate": 0.0,
            "bounce_rate": 0.0,
        }
        for filename in filenames
    }

    for row in sent_rows:
        filename = row["campaign_filename"]
        if filename in result:
            result[filename]["sent"] = int(row["sent"] or 0)

    for row in state_rows:
        filename = row["campaign_filename"]
        status = row["status"] or ""
        total = int(row["total"] or 0)

        if filename not in result:
            continue

        if status == "replied":
            result[filename]["replies"] = total
        elif status == "bounced":
            result[filename]["bounces"] = total
        elif status == "unsubscribed":
            result[filename]["unsubscribes"] = total

    for filename, data in result.items():
        data["reply_rate"] = _pct(data["replies"], data["sent"])
        data["bounce_rate"] = _pct(data["bounces"], data["sent"])

    return result


@router.get("/api/campaigns/{campaign_filename}/runs")
def get_campaign_runs(campaign_filename: str) -> list[dict]:
    runs = _campaign_runs(campaign_filename)
    return [
        {
            "id": run.id,
            "label": _run_label(run),
            "status": run.status.value,
            "started_at": _dt(run.started_at) or "",
            "completed_at": _dt(getattr(run, "completed_at", None)),
            "total_scraped": run.total_scraped,
            "total_warm": run.total_warm,
            "total_cold": run.total_cold,
            "total_no_email": run.total_no_email,
            "start_url": (run.filters or {}).get("start_url", ""),
            "campaign": (
                (run.filters or {}).get("campaign")
                or (run.filters or {}).get("campaign_key")
                or ""
            ),
        }
        for run in runs
    ]

@router.get("/api/campaigns/{campaign_filename}/overview")
def get_campaign_overview(campaign_filename: str) -> dict:
    runs = get_campaign_runs(campaign_filename)
    rows = _campaign_lead_rows(
        campaign_filename,
        limit=None,
    )
    total_leads = len(rows)
    with_email = sum(1 for row in rows if row.get("email"))
    no_email = total_leads - with_email

    coverage = lead_universe_repo.campaign_coverage(campaign_filename)
    coverage.update({
        "needs_enrichment": no_email,
        "with_email": with_email,
    })

    counts = outreach_repo.campaign_overview_counts(campaign_filename)
    draft_counts = counts["draft_counts"]
    drafts_generated = counts["drafts_generated"]
    drafted_unique = counts["drafted_unique"]
    approved_unique = counts["approved_unique"]
    sent_unique = counts["sent_unique"]
    state_counts = counts["state_counts"]
    active_sequence_steps = len(_active_steps(campaign_filename))
    approved_drafts = draft_counts.get("approved", 0)
    scheduled = (
        draft_counts.get("scheduled", 0)
        + draft_counts.get("approved", 0)
    )
    emails_sent = draft_counts.get("sent", 0)
    followups_due = len(outreach_repo.due_items(campaign_filename))
    replies = state_counts.get("replied", 0)
    bounces = state_counts.get("bounced", 0)
    unsubscribed = state_counts.get("unsubscribed", 0)
    completed = state_counts.get("completed", 0)

    return {
        "campaign_filename": campaign_filename,
        "total_leads": total_leads,
        "needs_enrichment": no_email,
        "with_email": with_email,
        "no_email": no_email,
        "drafts_generated": drafts_generated,
        "approved_drafts": approved_drafts,
        "scheduled": scheduled,
        "emails_sent": emails_sent,
        "followups_due": followups_due,
        "replies": replies,
        "bounces": bounces,
        "unsubscribed": unsubscribed,
        "completed": completed,
        "active_sequence_steps": active_sequence_steps,
        "pipeline": {
            "scraped": total_leads,
            "enriched": with_email,
            "drafted": drafted_unique,
            "approved": approved_unique,
            "sent": sent_unique,
            "replied": replies,
            "completed": completed,
        },
        "total_runs": len(runs),
        "runs": runs,
        "lead_collection": coverage,
    }



@router.get("/api/campaigns/{campaign_filename}/report")
def get_campaign_report(
    campaign_filename: str,
    days: int = Query(default=30, ge=1, le=365),
) -> dict:
    return _campaign_report_data(campaign_filename, days)


@router.get("/api/campaigns/{campaign_filename}/leads")
def get_campaign_leads(
    campaign_filename: str,
    segment: Optional[str] = None,
    run_id: Optional[str] = None,
    limit: int = Query(default=500, ge=1, le=2000),
    offset: int = Query(default=0, ge=0),
    page: Optional[int] = None,
    page_size: int = Query(default=50, ge=1, le=200),
    q: str = "",
    sequence_status: str = "",
) -> dict | list[dict]:
    if page is None:
        rows = _campaign_lead_rows(
            campaign_filename,
            segment=segment,
            run_id=run_id,
            limit=limit,
            offset=offset,
        )
        payloads = [_campaign_lead_payload(row) for row in rows]
        _attach_other_campaigns(payloads, campaign_filename)
        return payloads

    page = max(1, int(page))
    page_size = min(200, max(1, int(page_size)))

    items, total = lead_repo.search_campaign_page(
        campaign_filename=campaign_filename,
        q=q.strip(),
        segment=segment or "",
        sequence_status=sequence_status or "",
        limit=page_size,
        offset=(page - 1) * page_size,
    )

    payloads = [_campaign_lead_payload(row) for row in items]
    _attach_other_campaigns(payloads, campaign_filename)

    return {
        "items": payloads,
        "total": total,
        "page": page,
        "page_size": page_size,
    }

@router.get("/api/campaigns/{campaign_filename}/activities")
def get_campaign_activities(
    campaign_filename: str,
    limit: int = Query(default=100, ge=1, le=500),
) -> list[dict]:
    return [
        _activity_payload(row)
        for row in outreach_repo.list_campaign_activities(
            campaign_filename,
            limit=limit,
        )
    ]

@router.get("/api/campaigns/{campaign_filename}/export-zoominfo")
def export_campaign_for_zoominfo(campaign_filename: str):
    from fastapi.responses import StreamingResponse

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Lead ID",
        "Run ID",
        "First Name",
        "Last Name",
        "Full Name",
        "Company Name",
        "Job Title",
        "LinkedIn URL",
        "Company LinkedIn URL",
        "Location",
        "Email",
        "Phone",
    ])
    safe_name = campaign_filename.replace(".json", "")
    for lead in _campaign_leads(campaign_filename):
        first_name = lead.first_name
        last_name = lead.last_name
        if not first_name and lead.full_name:
            parts = lead.full_name.split()
            first_name = parts[0]
            last_name = " ".join(parts[1:])
        writer.writerow([
            lead.id,
            getattr(lead, "run_id", "") or "",
            first_name,
            last_name,
            lead.full_name,
            lead.company,
            lead.title,
            lead.linkedin_url,
            lead.company_linkedin_url,
            lead.location,
            lead.email,
            lead.phone,
        ])
        _add_activity(
            lead,
            campaign_filename,
            "exported_for_zoominfo",
            "Lead exported for ZoomInfo",
            "Campaign ZoomInfo export downloaded",
            {"filename": f"{safe_name}_zoominfo_export.csv"},
        )
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={
            "Content-Disposition": (
                f"attachment; filename={safe_name}_zoominfo_export.csv"
            )
        },
    )
def _reconciliation_candidate_payload(lead: Lead) -> dict:
    return {
        "id": lead.id,
        "full_name": lead.full_name or "",
        "company": lead.company or "",
        "title": lead.title or "",
        "email": lead.email or "",
        "linkedin_url": lead.linkedin_url or "",
    }


def _reconciliation_row_label(row: dict, row_number: int) -> str:
    full_name = _row_value(row, "Full Name", "Name", "full_name")
    email = _row_value(
        row,
        "Email",
        "Work Email",
        "Business Email",
        "Email Address",
        "email",
    )
    company = _row_value(
        row,
        "Company Name",
        "Company",
        "Account Name",
        "company",
    )

    label = " · ".join(part for part in [full_name, email, company] if part)
    return label or f"Row {row_number}"


def _add_reconciliation_diff(
    diffs: list[dict],
    field: str,
    current_value,
    new_value,
    only_if_empty: bool = False,
) -> None:
    old = "" if current_value is None else str(current_value).strip()
    new = "" if new_value is None else str(new_value).strip()

    if not new:
        return

    if only_if_empty and old:
        return

    if old != new:
        diffs.append(
            {
                "field": field,
                "old": old,
                "new": new,
            }
        )


async def _read_enriched_rows(file: UploadFile) -> list[dict]:
    contents = await file.read()
    filename = (file.filename or "").lower()

    if filename.endswith(".xlsx"):
        try:
            import openpyxl
        except ImportError as exc:
            raise HTTPException(
                status_code=500,
                detail="openpyxl not installed. Run: pip install openpyxl",
            ) from exc

        workbook = openpyxl.load_workbook(
            io.BytesIO(contents),
            read_only=True,
            data_only=True,
        )
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(cell or "").strip() for cell in rows[0]]
        parsed = []
        for values in rows[1:]:
            parsed.append({
                headers[idx]: "" if value is None else str(value).strip()
                for idx, value in enumerate(values)
                if idx < len(headers) and headers[idx]
            })
        return parsed

    if filename.endswith(".csv") or not filename:
        text = contents.decode("utf-8-sig", errors="ignore")
        return list(csv.DictReader(io.StringIO(text)))

    raise HTTPException(
        status_code=400,
        detail="Only .csv and .xlsx files are supported",
    )


def _uploaded_row_to_lead(row: dict) -> Lead | None:
    email = _row_value(
        row,
        "Email",
        "Work Email",
        "Business Email",
        "Email Address",
        "email",
    )
    full_name = _row_value(row, "Full Name", "Name", "full_name")
    first_name = _row_value(row, "First Name", "first_name")
    last_name = _row_value(row, "Last Name", "last_name")

    if not full_name:
        full_name = " ".join(part for part in [first_name, last_name] if part).strip()

    if not first_name and full_name:
        parts = full_name.split()
        first_name = parts[0] if parts else ""
        last_name = " ".join(parts[1:]) if len(parts) > 1 else ""

    company = _row_value(
        row,
        "Company Name",
        "Company",
        "Account Name",
        "company",
    )
    title = _row_value(row, "Job Title", "Title", "title")
    phone = _row_value(
        row,
        "Phone",
        "Direct Phone",
        "Mobile Phone",
        "Cell",
        "Phone Number",
        "phone",
    )
    domain = _row_value(
        row,
        "Company Domain",
        "Domain",
        "Website",
        "Company Website",
        "company_domain",
    )
    linkedin_url = _norm_url(
        _row_value(
            row,
            "LinkedIn URL",
            "Person LinkedIn URL",
            "linkedin_url",
        )
    )
    company_linkedin_url = _norm_url(
        _row_value(
            row,
            "Company LinkedIn URL",
            "company_linkedin_url",
        )
    )
    location = _split_location(row)

    if domain:
        import re
        domain = re.sub(
            r"https?://(www\.)?",
            "",
            domain,
            flags=re.IGNORECASE,
        ).rstrip("/")

    # Require at least one useful identity field.
    if not any([email, linkedin_url, full_name, company]):
        return None

    return Lead(
        full_name=full_name,
        first_name=first_name,
        last_name=last_name,
        title=title,
        company=company,
        company_domain=domain,
        location=location,
        linkedin_url=linkedin_url,
        company_linkedin_url=company_linkedin_url,
        email=email,
        email_confidence="uploaded" if email else "",
        phone=phone,
        segment=Segment.WARM if email else Segment.NO_EMAIL,
        status=LeadStatus.ENRICHED if email else LeadStatus.SCRAPED,
    )


def _index_reconciliation_lead(
    lead: Lead,
    by_id,
    by_linkedin,
    by_email,
    by_name_company,
    by_first_last_company,
) -> None:
    by_id[lead.id] = lead
    if lead.linkedin_url:
        by_linkedin[_norm_url(lead.linkedin_url)].append(lead)
    if lead.email:
        by_email[lead.email.strip().lower()].append(lead)
    if lead.full_name and lead.company:
        by_name_company[_norm_name_company(lead.full_name, lead.company)].append(lead)
    if lead.first_name and lead.last_name and lead.company:
        by_first_last_company[
            _norm_name_company(f"{lead.first_name} {lead.last_name}", lead.company)
        ].append(lead)


        
@router.post("/api/campaigns/{campaign_filename}/upload-enriched")
async def upload_campaign_enriched(
    campaign_filename: str,
    file: UploadFile = File(...),
) -> dict:
    rows = await _read_enriched_rows(file)
    total_rows = len(rows)

    report = {
        "updated": [],
        "unchanged": [],
        "unmatched": [],
        "ambiguous": [],
        "errors": [],
    }

    if not rows:
        return {
            "success": True,
            "total_rows": 0,
            "matched": 0,
            "updated": 0,
            "unchanged": 0,
            "unmatched": 0,
            "ambiguous": 0,
            "errors": [],
            "report": report,
            "message": "No rows found in uploaded file",
        }

    run_ids = {
        _row_value(row, "Run ID", "run_id")
        for row in rows
        if _row_value(row, "Run ID", "run_id")
    }
    campaign_run_ids = set(_campaign_run_ids(campaign_filename))
    if run_ids:
        run_ids = {run_id for run_id in run_ids if run_id in campaign_run_ids}

    if run_ids:
        leads = []
        for run_id in run_ids:
            leads.extend(lead_repo.get_by_run(run_id))
    else:
        leads = _campaign_leads(campaign_filename)

    from collections import defaultdict

    by_id = {lead.id: lead for lead in leads}
    by_linkedin: dict[str, list[Lead]] = defaultdict(list)
    by_email: dict[str, list[Lead]] = defaultdict(list)
    by_name_company: dict[str, list[Lead]] = defaultdict(list)
    by_first_last_company: dict[str, list[Lead]] = defaultdict(list)

    for lead in leads:
        if lead.linkedin_url:
            by_linkedin[_norm_url(lead.linkedin_url)].append(lead)
        if lead.email:
            by_email[lead.email.strip().lower()].append(lead)
        if lead.full_name and lead.company:
            by_name_company[_norm_name_company(lead.full_name, lead.company)].append(lead)
        if lead.first_name and lead.last_name and lead.company:
            by_first_last_company[
                _norm_name_company(f"{lead.first_name} {lead.last_name}", lead.company)
            ].append(lead)

    matched = 0
    updated = 0
    unchanged = 0
    unmatched = 0
    ambiguous = 0
    created = 0
    created_with_email = 0
    affected_runs: set[str] = set()
    manual_upload_run = None

    def ensure_manual_upload_run() -> PipelineRun:
        nonlocal manual_upload_run
        if manual_upload_run:
            return manual_upload_run

        manual_upload_run = PipelineRun(
            status=RunStatus.RUNNING,
            filters={
                "campaign": campaign_filename,
                "campaign_key": campaign_filename,
                "source": "manual_upload",
                "upload_filename": file.filename or "",
            },
            enrichment_mode=EnrichmentMode.FREE,
            started_at=datetime.utcnow(),
        )
        run_repo.save(manual_upload_run)
        return manual_upload_run

    for idx, row in enumerate(rows, start=2):
        row_label = _reconciliation_row_label(row, idx)

        lead_id = _row_value(row, "Lead ID", "lead_id", "id")
        linkedin = _norm_url(
            _row_value(
                row,
                "LinkedIn URL",
                "Person LinkedIn URL",
                "linkedin_url",
            )
        )
        email = _row_value(
            row,
            "Email",
            "Work Email",
            "Business Email",
            "Email Address",
            "email",
        )
        full_name = _row_value(row, "Full Name", "Name", "full_name")
        company = _row_value(
            row,
            "Company Name",
            "Company",
            "Account Name",
            "company",
        )
        first_name = _row_value(row, "First Name", "first_name")
        last_name = _row_value(row, "Last Name", "last_name")

        lead = None
        match_method = ""

        if lead_id:
            lead = by_id.get(lead_id)
            if lead:
                match_method = "lead_id"

        if not lead and linkedin:
            candidates = by_linkedin.get(linkedin, [])
            if len(candidates) == 1:
                lead = candidates[0]
                match_method = "linkedin_url"
            elif len(candidates) > 1:
                ambiguous += 1
                report["ambiguous"].append(
                    {
                        "row": idx,
                        "label": row_label,
                        "reason": "multiple_linkedin_matches",
                        "match_method": "linkedin_url",
                        "candidates": [
                            _reconciliation_candidate_payload(candidate)
                            for candidate in candidates
                        ],
                        "input": row,
                    }
                )
                continue

        if not lead and email:
            candidates = by_email.get(email.strip().lower(), [])
            if len(candidates) == 1:
                lead = candidates[0]
                match_method = "email"
            elif len(candidates) > 1:
                ambiguous += 1
                report["ambiguous"].append(
                    {
                        "row": idx,
                        "label": row_label,
                        "reason": "multiple_email_matches",
                        "match_method": "email",
                        "candidates": [
                            _reconciliation_candidate_payload(candidate)
                            for candidate in candidates
                        ],
                        "input": row,
                    }
                )
                continue

        if not lead and full_name and company:
            candidates = by_name_company.get(_norm_name_company(full_name, company), [])
            if len(candidates) == 1:
                lead = candidates[0]
                match_method = "full_name_company"
            elif len(candidates) > 1:
                ambiguous += 1
                report["ambiguous"].append(
                    {
                        "row": idx,
                        "label": row_label,
                        "reason": "multiple_name_company_matches",
                        "match_method": "full_name_company",
                        "candidates": [
                            _reconciliation_candidate_payload(candidate)
                            for candidate in candidates
                        ],
                        "input": row,
                    }
                )
                continue

        if not lead and first_name and last_name and company:
            candidates = by_first_last_company.get(
                _norm_name_company(f"{first_name} {last_name}", company),
                [],
            )
            if len(candidates) == 1:
                lead = candidates[0]
                match_method = "first_last_company"
            elif len(candidates) > 1:
                ambiguous += 1
                report["ambiguous"].append(
                    {
                        "row": idx,
                        "label": row_label,
                        "reason": "multiple_first_last_company_matches",
                        "match_method": "first_last_company",
                        "candidates": [
                            _reconciliation_candidate_payload(candidate)
                            for candidate in candidates
                        ],
                        "input": row,
                    }
                )
                continue

        if not lead:
            new_lead = _uploaded_row_to_lead(row)
            if not new_lead:
                unmatched += 1
                report["unmatched"].append(
                    {
                        "row": idx,
                        "label": row_label,
                        "reason": "no_matching_lead_or_required_identity_fields",
                        "input": row,
                    }
                )
                continue

            upload_run = ensure_manual_upload_run()
            lead_repo.save_batch(upload_run.id, [new_lead])
            affected_runs.add(upload_run.id)
            created += 1
            if new_lead.email:
                created_with_email += 1

            verification = verify_email(new_lead.email) if new_lead.email else None
            if verification:
                lead_repo.set_email_verification(
                    new_lead.id,
                    verification["status"],
                    verification["reason"],
                    verification["checked_at"],
                )

            state = outreach_repo.get_or_create_state(
                new_lead.id,
                campaign_filename,
            )
            state.status = "enriched" if new_lead.email else "not_started"
            outreach_repo.upsert_state(state)
            _set_lead_sequence_columns(new_lead.id, state.status)

            lead_repo.mark_duplicate_if_any(new_lead.id, campaign_filename)
            refreshed = lead_repo.get_by_id(new_lead.id) or new_lead

            _add_activity(
                refreshed,
                campaign_filename,
                "lead_imported",
                "Lead imported from uploaded file",
                "This lead was created from a campaign upload.",
                {
                    "row": idx,
                    "source": "manual_upload",
                    "email_verification": verification,
                },
            )

            _index_reconciliation_lead(
                refreshed,
                by_id,
                by_linkedin,
                by_email,
                by_name_company,
                by_first_last_company,
            )

            report["updated"].append(
                {
                    "row": idx,
                    "label": row_label,
                    "lead": _reconciliation_candidate_payload(refreshed),
                    "match_method": "created_from_upload",
                    "diffs": [],
                    "email_verification": verification,
                    "input": row,
                }
            )
            continue

        matched += 1

        email_value = email
        phone = _row_value(
            row,
            "Phone",
            "Direct Phone",
            "Mobile Phone",
            "Cell",
            "Phone Number",
            "phone",
        )
        domain = _row_value(
            row,
            "Company Domain",
            "Domain",
            "Website",
            "Company Website",
            "company_domain",
        )
        title = _row_value(row, "Job Title", "Title", "title")
        location = _split_location(row)
        linkedin_url = _norm_url(
            _row_value(
                row,
                "LinkedIn URL",
                "Person LinkedIn URL",
                "linkedin_url",
            )
        )
        company_linkedin_url = _norm_url(
            _row_value(
                row,
                "Company LinkedIn URL",
                "company_linkedin_url",
            )
        )

        if domain:
            import re

            domain = re.sub(
                r"https?://(www\.)?",
                "",
                domain,
                flags=re.IGNORECASE,
            ).rstrip("/")

        updates: dict[str, str] = {}
        diffs: list[dict] = []

        if email_value:
            updates["email"] = email_value
            _add_reconciliation_diff(diffs, "email", lead.email, email_value)

        if phone:
            updates["phone"] = phone
            _add_reconciliation_diff(diffs, "phone", lead.phone, phone)

        if domain:
            updates["company_domain"] = domain
            _add_reconciliation_diff(
                diffs,
                "company_domain",
                getattr(lead, "company_domain", ""),
                domain,
            )

        if title:
            updates["title"] = title
            _add_reconciliation_diff(
                diffs,
                "title",
                lead.title,
                title,
                only_if_empty=True,
            )

        if location:
            updates["location"] = location
            _add_reconciliation_diff(
                diffs,
                "location",
                lead.location,
                location,
                only_if_empty=True,
            )

        if linkedin_url:
            updates["linkedin_url"] = linkedin_url
            _add_reconciliation_diff(
                diffs,
                "linkedin_url",
                lead.linkedin_url,
                linkedin_url,
                only_if_empty=True,
            )

        if company_linkedin_url:
            updates["company_linkedin_url"] = company_linkedin_url
            _add_reconciliation_diff(
                diffs,
                "company_linkedin_url",
                getattr(lead, "company_linkedin_url", ""),
                company_linkedin_url,
                only_if_empty=True,
            )

        verification = None
        if email_value:
            verification = verify_email(email_value)

        if not updates or not diffs:
            unchanged += 1

            if verification:
                lead_repo.set_email_verification(
                    lead.id,
                    verification["status"],
                    verification["reason"],
                    verification["checked_at"],
                )

            report["unchanged"].append(
                {
                    "row": idx,
                    "label": row_label,
                    "lead": _reconciliation_candidate_payload(lead),
                    "match_method": match_method,
                    "reason": "matched_but_no_new_values",
                    "email_verification": verification,
                    "input": row,
                }
            )
            continue

        try:
            changed = lead_repo.update_lead_enrichment(lead.id, updates)

            if verification:
                lead_repo.set_email_verification(
                    lead.id,
                    verification["status"],
                    verification["reason"],
                    verification["checked_at"],
                )

            if changed:
                updated += 1
                affected_runs.add(lead.run_id)
                refreshed = lead_repo.get_by_id(lead.id) or lead

                if email_value:
                    state = outreach_repo.get_or_create_state(
                        lead.id,
                        campaign_filename,
                    )
                    if state.status in {"not_started", "enriched"}:
                        state.status = "enriched"
                        outreach_repo.upsert_state(state)
                        _set_lead_sequence_columns(lead.id, "enriched")

                    if not (lead.email or "").strip():
                        lead_repo.mark_duplicate_if_any(
                            lead.id,
                            campaign_filename,
                        )

                _add_activity(
                    refreshed,
                    campaign_filename,
                    "enriched",
                    "Lead enriched from uploaded file",
                    "ZoomInfo enrichment data matched and updated this lead",
                    {
                        "fields": [diff["field"] for diff in diffs],
                        "diffs": diffs,
                        "row": idx,
                        "email_verification": verification,
                    },
                )

                report["updated"].append(
                    {
                        "row": idx,
                        "label": row_label,
                        "lead": _reconciliation_candidate_payload(refreshed),
                        "match_method": match_method,
                        "diffs": diffs,
                        "email_verification": verification,
                        "input": row,
                    }
                )
            else:
                unchanged += 1
                report["unchanged"].append(
                    {
                        "row": idx,
                        "label": row_label,
                        "lead": _reconciliation_candidate_payload(lead),
                        "match_method": match_method,
                        "reason": "repository_reported_no_change",
                        "email_verification": verification,
                        "input": row,
                    }
                )
        except Exception as exc:
            
            report["errors"].append(
                {
                    "row": idx,
                    "label": row_label,
                    "reason": str(exc),
                    "input": row,
                }
            )

    if manual_upload_run:
        manual_upload_run.status = RunStatus.COMPLETED
        manual_upload_run.total_scraped = created
        manual_upload_run.total_enriched = created_with_email
        manual_upload_run.total_warm = created_with_email
        manual_upload_run.total_no_email = max(0, created - created_with_email)
        manual_upload_run.completed_at = datetime.utcnow()
        run_repo.save(manual_upload_run)

    _update_segments_for_runs({run_id for run_id in affected_runs if run_id})

    errors = [
        f"Row {item['row']}: {item['reason']}"
        for item in report["errors"]
    ]

    return {
        "success": True,
        "total_rows": total_rows,
        "matched": matched,
        "updated": updated,
        "unchanged": unchanged,
        "unmatched": unmatched,
        "ambiguous": ambiguous,
        "created": created,
        "errors": errors,
        "report": report,
        "message": (
            "No usable leads found. Make sure the file includes Email, LinkedIn URL, Full Name, or Company."
            if matched == 0 and created == 0 and total_rows > 0
            else f"Upload processed successfully. Created {created}, updated {updated}."
        ),
    }


@router.post("/api/campaigns/{campaign_filename}/verify-emails")
def verify_campaign_emails(
    campaign_filename: str,
    only_missing: bool = Query(True),
    limit: int = Query(500, ge=1, le=5000),
) -> dict:
    rows = _campaign_lead_rows(campaign_filename, limit=None)

    checked = 0
    skipped = 0
    counts = {
        "valid": 0,
        "risky": 0,
        "invalid": 0,
        "missing": 0,
    }
    results = []

    for row in rows:
        if checked >= limit:
            break

        lead_id = str(row.get("id") or "").strip()
        email = str(row.get("email") or "").strip()
        current_status = str(row.get("email_verification_status") or "").strip()

        if not lead_id:
            skipped += 1
            continue

        if not email:
            counts["missing"] += 1
            skipped += 1
            results.append({
                "lead_id": lead_id,
                "email": "",
                "status": "missing",
                "reason": "no_email",
            })
            continue

        if only_missing and current_status:
            skipped += 1
            counts[current_status] = counts.get(current_status, 0) + 1
            continue

        verification = verify_email(email)
        lead_repo.set_email_verification(
            lead_id,
            verification["status"],
            verification["reason"],
            verification["checked_at"],
        )

        checked += 1
        status = verification["status"]
        counts[status] = counts.get(status, 0) + 1
        results.append({
            "lead_id": lead_id,
            "email": email,
            "status": status,
            "reason": verification["reason"],
            "checked_at": verification["checked_at"],
        })

    return {
        "success": True,
        "campaign_filename": campaign_filename,
        "checked": checked,
        "skipped": skipped,
        "counts": counts,
        "only_missing": only_missing,
        "limit": limit,
        "results": results,
        "message": f"Verified {checked} emails.",
    }


@router.get("/api/campaigns/{campaign_filename}/sequence-settings")
def get_campaign_sequence_settings(campaign_filename: str) -> dict:
    return _load_sequence_settings(campaign_filename)

@router.post("/api/campaigns/{campaign_filename}/sequence-settings")
def save_campaign_sequence_settings(
    campaign_filename: str,
    request: SequenceSettingsRequest,
) -> dict:
    touches = []
    incoming = request.steps or request.touches
    if not incoming:
        raise HTTPException(
            status_code=400,
            detail="At least one sequence step is required",
        )
    for touch in incoming:
        if touch.number <= 0:
            raise HTTPException(
                status_code=400,
                detail="touch_number must be positive",
            )
        if touch.delay_days < 0:
            raise HTTPException(
                status_code=400,
                detail="delay_days cannot be negative",
            )
        if hasattr(touch, "model_dump"):
            item = touch.model_dump()
        else:
            item = touch.dict()
        item["touch_number"] = item.get("touch_number") or item.get("number")
        item["touch_name"] = item.get("touch_name") or item.get("name")
        touches.append(item)
    rules = request.rules.model_dump() if request.rules else {}
    _save_sequence_settings(
        campaign_filename,
        {
            "name": request.name,
            "sequence_name": request.sequence_name,
            "touches": touches,
            "steps": touches,
            "rules": rules,
        },
    )
    return {
        "saved": True,
        "campaign": campaign_filename,
        **_load_sequence_settings(campaign_filename),
    }

@router.post("/api/campaigns")
def create_campaign(request: CreateCampaignRequest) -> dict:
    """Create a new campaign in the local SQLite database."""
    if not request.name.strip():
        raise HTTPException(status_code=400, detail="Campaign name is required")

    config = request.model_dump(exclude={"name", "description"})
    try:
        campaign = campaign_repo.create(
            request.name.strip(),
            request.description,
            config,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    return {
        "created": True,
        "filename": campaign["filename"],
        "name": campaign["name"],
    }

@router.patch("/api/campaigns/{campaign_filename}")
def update_campaign(
    campaign_filename: str,
    request: CampaignUpdateRequest,
) -> dict:
    campaign = campaign_repo.get_by_filename(campaign_filename)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")

    updates = request.model_dump(exclude_none=True)
    cleaned = {
        key: str(value or "").strip()
        for key, value in updates.items()
    }

    config = {
        **(campaign.get("config") or {}),
        **cleaned,
    }

    updated = campaign_repo.update_config(campaign_filename, config)
    if not updated:
        raise HTTPException(status_code=404, detail="Campaign not found")

    return updated


def _campaign_filename_aliases(campaign_filename: str) -> set[str]:
    normalized = (campaign_filename or "").strip()
    if not normalized:
        return set()
    with_json = normalized if normalized.endswith(".json") else f"{normalized}.json"
    without_json = with_json[:-5] if with_json.endswith(".json") else with_json
    return {with_json, without_json}


def _delete_by_ids(conn, table: str, column: str, ids: set[str]) -> int:
    values = sorted(str(value) for value in ids if value)
    if not values:
        return 0
    placeholders = ",".join("?" for _ in values)
    cur = conn.execute(
        f"DELETE FROM {table} WHERE {column} IN ({placeholders})",
        values,
    )
    return int(cur.rowcount or 0)


def _job_payload_campaign(payload_json: str) -> str:
    try:
        payload = json.loads(payload_json or "{}")
    except Exception:
        return ""
    return str(
        payload.get("campaign_filename")
        or payload.get("campaign")
        or payload.get("campaign_key")
        or ""
    ).strip()


def _run_filters_campaign(filters_json: str) -> str:
    try:
        filters = json.loads(filters_json or "{}")
    except Exception:
        return ""
    return str(
        filters.get("campaign")
        or filters.get("campaign_key")
        or filters.get("campaign_filename")
        or ""
    ).strip()


@router.delete("/api/campaigns/{campaign_filename}")
def delete_campaign(campaign_filename: str) -> dict:
    campaign = campaign_repo.get_by_filename(campaign_filename)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")

    filename = campaign["filename"]
    aliases = _campaign_filename_aliases(filename)
    conn = lead_repo.db.conn()

    active_bulk_rows = conn.execute(
        """
        SELECT id
        FROM bulk_scrape_jobs
        WHERE campaign_key IN (?, ?)
          AND status IN ('queued', 'running')
        """,
        tuple(sorted(aliases)),
    ).fetchall()

    if active_bulk_rows:
        raise HTTPException(
            status_code=409,
            detail="Stop or cancel active bulk scrape jobs before deleting this campaign.",
        )

    active_run_rows = conn.execute(
        """
        SELECT id, filters
        FROM pipeline_runs
        WHERE status = 'RUNNING'
        """
    ).fetchall()
    for row in active_run_rows:
        if _run_filters_campaign(row["filters"]) in aliases:
            raise HTTPException(
                status_code=409,
                detail="Stop active scrape runs before deleting this campaign.",
            )

    active_job_rows = conn.execute(
        """
        SELECT id, payload_json
        FROM jobs
        WHERE status IN ('queued', 'running')
        """
    ).fetchall()
    for row in active_job_rows:
        if _job_payload_campaign(row["payload_json"]) in aliases:
            raise HTTPException(
                status_code=409,
                detail="Cancel active campaign jobs before deleting this campaign.",
            )

    bulk_rows = conn.execute(
        """
        SELECT id
        FROM bulk_scrape_jobs
        WHERE campaign_key IN (?, ?)
        """,
        tuple(sorted(aliases)),
    ).fetchall()
    bulk_job_ids = {row["id"] for row in bulk_rows}

    run_rows = conn.execute(
        """
        SELECT id, filters
        FROM pipeline_runs
        """
    ).fetchall()

    run_ids: set[str] = set()
    for row in run_rows:
        try:
            filters = json.loads(row["filters"] or "{}")
        except Exception:
            filters = {}
        campaign_key = str(
            filters.get("campaign")
            or filters.get("campaign_key")
            or filters.get("campaign_filename")
            or ""
        ).strip()
        bulk_job_id = str(filters.get("bulk_scrape_job_id") or "").strip()
        if campaign_key in aliases or bulk_job_id in bulk_job_ids:
            run_ids.add(row["id"])

    universe_rows = conn.execute(
        """
        SELECT id
        FROM lead_universes
        WHERE campaign_filename = ?
        """,
        (filename,),
    ).fetchall()
    universe_ids = {row["id"] for row in universe_rows}

    segment_rows = conn.execute(
        """
        SELECT id
        FROM lead_source_segments
        WHERE campaign_filename = ?
        """,
        (filename,),
    ).fetchall()
    segment_ids = {row["id"] for row in segment_rows}

    if universe_ids:
        placeholders = ",".join("?" for _ in universe_ids)
        rows = conn.execute(
            f"""
            SELECT id
            FROM lead_source_segments
            WHERE universe_id IN ({placeholders})
            """,
            sorted(universe_ids),
        ).fetchall()
        segment_ids.update(row["id"] for row in rows)

    job_rows = conn.execute(
        """
        SELECT id, payload_json
        FROM jobs
        """
    ).fetchall()
    job_ids = {
        row["id"]
        for row in job_rows
        if _job_payload_campaign(row["payload_json"]) in aliases
    }

    counts = {}

    with conn:
        counts["outreach_drafts"] = conn.execute(
            "DELETE FROM outreach_drafts WHERE campaign_filename = ?",
            (filename,),
        ).rowcount
        counts["lead_sequence_state"] = conn.execute(
            "DELETE FROM lead_sequence_state WHERE campaign_filename = ?",
            (filename,),
        ).rowcount
        counts["lead_activities"] = conn.execute(
            "DELETE FROM lead_activities WHERE campaign_filename = ?",
            (filename,),
        ).rowcount
        counts["send_log"] = conn.execute(
            "DELETE FROM send_log WHERE campaign_filename = ?",
            (filename,),
        ).rowcount
        counts["campaign_sequence_steps"] = conn.execute(
            "DELETE FROM campaign_sequence_steps WHERE campaign_filename IN (?, ?)",
            tuple(sorted(aliases)),
        ).rowcount
        counts["campaign_sequence_rules"] = conn.execute(
            "DELETE FROM campaign_sequence_rules WHERE campaign_filename IN (?, ?)",
            tuple(sorted(aliases)),
        ).rowcount

        counts["run_checkpoints"] = _delete_by_ids(conn, "run_checkpoints", "run_id", run_ids)

        try:
            counts["run_controls"] = _delete_by_ids(conn, "run_controls", "run_id", run_ids)
        except Exception:
            counts["run_controls"] = 0

        counts["leads_by_run"] = _delete_by_ids(conn, "leads", "run_id", run_ids)
        counts["leads_by_universe"] = _delete_by_ids(conn, "leads", "lead_universe_id", universe_ids)
        counts["leads_by_segment"] = _delete_by_ids(conn, "leads", "lead_source_segment_id", segment_ids)

        counts["lead_source_segments_by_id"] = _delete_by_ids(conn, "lead_source_segments", "id", segment_ids)
        counts["lead_source_segments_by_campaign"] = conn.execute(
            "DELETE FROM lead_source_segments WHERE campaign_filename = ?",
            (filename,),
        ).rowcount
        counts["lead_universes"] = conn.execute(
            "DELETE FROM lead_universes WHERE campaign_filename = ?",
            (filename,),
        ).rowcount

        counts["bulk_scrape_jobs"] = _delete_by_ids(conn, "bulk_scrape_jobs", "id", bulk_job_ids)
        counts["pipeline_runs"] = _delete_by_ids(conn, "pipeline_runs", "id", run_ids)
        counts["jobs"] = _delete_by_ids(conn, "jobs", "id", job_ids)

        counts["campaigns"] = conn.execute(
            "DELETE FROM campaigns WHERE filename = ?",
            (filename,),
        ).rowcount

    return {
        "deleted": True,
        "filename": filename,
        "counts": counts,
    }



@router.get("/api/campaigns/{campaign_filename}/sequence/sample-leads")
def sequence_sample_leads(
    campaign_filename: str,
    q: str = "",
    limit: int = Query(default=50, ge=1, le=200),
) -> list[dict]:
    # Light, name-first list for the sample email picker.
    items, _total = lead_repo.search_campaign_page(
        campaign_filename=campaign_filename,
        q=(q or "").strip(),
        segment="",
        sequence_status="",
        limit=limit,
        offset=0,
    )
    return [
        {
            "id": row.get("id", ""),
            "full_name": row.get("full_name", "") or "Unknown lead",
            "company": row.get("company", "") or "",
            "title": row.get("title", "") or "",
            "has_email": bool(row.get("email")),
        }
        for row in items
    ]


@router.post("/api/campaigns/{campaign_filename}/sequence/preview")
def sequence_preview(
    campaign_filename: str,
    request: SequencePreviewRequest,
) -> dict:
    # Generates a sample only. Persists nothing: no draft, no state, no send.
    subject_direction = (request.subject_template or "").strip()
    body_direction = (request.email_body_template or "").strip()
    if not subject_direction and not body_direction:
        raise HTTPException(
            status_code=400,
            detail="Add a subject direction or AI instructions before previewing.",
        )

    sample = None
    if request.sample_lead_id:
        sample = lead_repo.get_by_id(request.sample_lead_id)
    if sample is None:
        campaign_leads = _campaign_leads(campaign_filename)
        for lead in campaign_leads:
            if lead.email:
                sample = lead
                break
        if sample is None and campaign_leads:
            sample = campaign_leads[0]
    if sample is None:
        sample = Lead(
            full_name="Jordan Avery",
            first_name="Jordan",
            last_name="Avery",
            title="VP Engineering",
            company="Northwind Logistics",
            location="Chicago, IL",
            email="jordan.avery@example.com",
        )

    used_ai = False
    error = ""
    subject = ""
    body = ""
    linkedin_message = ""
    research_summary = ""

    if _ai_personalised_drafts_enabled():
        try:
            campaign_config = KnowledgeBaseLoader.load_campaign(campaign_filename)
            research_agent = WebResearchAgent()
            context_agent = ContextAgent(campaign_config)
            writer_agent = WriterAgent(
                campaign_config,
                touch1_template={
                    "subject_template": subject_direction,
                    "email_body_template": body_direction,
                    "linkedin_message_template": "",
                },
            )
            # Role inference only: no website scrape, no browser launch.
            research = research_agent._infer_from_role(sample)
            context = context_agent.get_context(sample, research)
            message = writer_agent.write(sample, research, context)
            if message.error:
                raise RuntimeError(message.error)
            subject = (message.email_subject or "").strip()
            body = (message.email_body or "").strip()
            linkedin_message = (message.linkedin_message or "").strip()
            research_summary = (message.research_summary or "").strip()
            if subject or body:
                used_ai = True
        except Exception as exc:
            error = str(exc)

    if not used_ai:
        subject = render_template(subject_direction, sample, campaign_filename) or subject_direction
        body = render_template(body_direction, sample, campaign_filename) or body_direction

    body = _ensure_sender_signature(body, campaign_filename)

    return {
        "subject": subject,
        "body": body,
        "linkedin_message": linkedin_message,
        "research_summary": research_summary,
        "used_ai": used_ai,
        "error": error,
        "sample_lead": {
            "id": sample.id,
            "full_name": sample.full_name,
            "title": sample.title,
            "company": sample.company,
        },
    }


def _norm_campaign_key(value: str) -> str:
    return (value or "").replace(".json", "").replace("_", " ").lower().strip()


def _norm_linkedin(value: str) -> str:
    return (value or "").strip().split("?", 1)[0].rstrip("/").lower()


def _attach_other_campaigns(items: list[dict], campaign_filename: str) -> None:
    # Names every OTHER campaign each person appears in, matched by email or LinkedIn.
    for item in items:
        item.setdefault("other_campaigns", [])
    if not items:
        return
    emails = [item.get("email", "") for item in items]
    links = [item.get("linkedin_url", "") for item in items]
    matched = lead_repo.leads_matching(emails, links)
    if not matched:
        return

    run_to_campaign = {}
    for run in run_repo.list_all():
        filters = run.filters or {}
        run_to_campaign[run.id] = (
            filters.get("campaign_key") or filters.get("campaign") or ""
        )
    name_map = {}
    for campaign in campaign_repo.list_all():
        name_map[_norm_campaign_key(campaign.get("filename", ""))] = (
            campaign.get("name") or campaign.get("filename", "")
        )
    current = _norm_campaign_key(campaign_filename)

    by_email = {}
    by_link = {}
    for row in matched:
        cf = run_to_campaign.get(row.get("run_id", ""), "")
        if not cf or _norm_campaign_key(cf) == current:
            continue
        display = name_map.get(_norm_campaign_key(cf)) or _norm_campaign_key(cf)
        em = (row.get("email", "") or "").strip().lower()
        lk = _norm_linkedin(row.get("linkedin_url", ""))
        if em:
            by_email.setdefault(em, [])
            if display not in by_email[em]:
                by_email[em].append(display)
        if lk:
            by_link.setdefault(lk, [])
            if display not in by_link[lk]:
                by_link[lk].append(display)

    for item in items:
        em = (item.get("email", "") or "").strip().lower()
        lk = _norm_linkedin(item.get("linkedin_url", ""))
        names = []
        for name in by_email.get(em, []) + by_link.get(lk, []):
            if name not in names:
                names.append(name)
        item["other_campaigns"] = names


@router.get("/api/campaigns/{campaign_filename}/sequence/members")
def sequence_members(
    campaign_filename: str,
    q: str = "",
    limit: int = Query(default=200, ge=1, le=1000),
) -> dict:
    rows, total = outreach_repo.list_sequence_members(campaign_filename, q=q, limit=limit)
    _attach_other_campaigns(rows, campaign_filename)
    items = [
        {
            "lead_id": row.get("lead_id", ""),
            "full_name": row.get("full_name", "") or "Unknown lead",
            "company": row.get("company", "") or "",
            "title": row.get("title", "") or "",
            "current_touch": int(row.get("current_touch") or 0),
            "status": row.get("status", "") or "not_started",
            "next_touch_due_at": row.get("next_touch_due_at") or "",
            "other_campaigns": row.get("other_campaigns", []),
        }
        for row in rows
    ]
    return {"items": items, "total": total}


@router.post("/api/campaigns/{campaign_filename}/sequence/members/{lead_id}/remove")
def remove_sequence_member(
    campaign_filename: str,
    lead_id: str,
) -> dict:
    # This campaign only. Cancels any scheduled-but-unsent email for this lead here.
    _stop_sequence(
        lead_id,
        campaign_filename,
        "removed",
        "manually removed from sequence",
    )
    return {"removed": True, "lead_id": lead_id}


@router.get("/api/campaigns")
def list_campaigns() -> list[dict]:
    return campaign_repo.list_all()
